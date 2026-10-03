"""Test de bout en bout de la page « Générer une messe » : un chant structuré de la bibliothèque, ajouté
aux blocs, ressort dans le PowerPoint avec son refrain en gras et dans l'ordre chanté."""

import io
from pathlib import Path

from pptx import Presentation
from streamlit.testing.v1 import AppTest

from context.models import Chant, LectureLiturgique, MomentLiturgique, SectionChant, TypeSection, TypeLecture
from tools.db_handler import create_chant, init_db

APP = str(Path(__file__).resolve().parents[1] / "app.py")


def _chant_structure() -> int:
    init_db()
    return create_chant(Chant(
        titre="Chant de bout en bout",
        paroles="Refrain inventé de bout en bout\n\nPremier couplet inventé",
        moments=[MomentLiturgique.ENTREE],
        structure=[
            SectionChant("R", TypeSection.REFRAIN, ["Refrain inventé de bout en bout"]),
            SectionChant("1", TypeSection.COUPLET, ["Premier couplet inventé"]),
        ],
        ordre=["R", "1", "R"],
    ))


def _lignes_du_pptx(octets: bytes) -> list[tuple[str, bool]]:
    """(texte, gras) de chaque ligne du corps de chaque diapositive."""
    lignes = []
    for diapo in Presentation(io.BytesIO(octets)).slides:
        zones = [s for s in diapo.shapes if s.has_text_frame]
        for paragraphe in zones[-1].text_frame.paragraphs:
            texte = "".join(r.text for r in paragraphe.runs)
            if texte:
                lignes.append((texte, bool(paragraphe.runs[0].font.bold)))
    return lignes


def test_chant_structure_ajoute_a_la_messe_sort_en_gras_dans_l_ordre_chante():
    _chant_structure()
    at = AppTest.from_file(APP, default_timeout=30).run()
    at.session_state["aelf_data"] = {
        "informations": {"jour_liturgique_nom": "Jour de test"},
        "lectures": [LectureLiturgique(TypeLecture.EVANGILE, "Jn 1", "Titre", "Intro", "Contenu inventé")],
    }
    at.run()
    assert not at.exception

    selecteur = next(s for s in at.selectbox if s.key == "chant_select")
    selecteur.set_value(selecteur.options.index("Chant de bout en bout"))
    next(b for b in at.button if b.key == "add_chant").click().run()
    assert not at.exception
    assert any(b.get("type") == "chant" for b in at.session_state["blocs"])

    next(b for b in at.button if b.label == "📥 Générer et télécharger PPTX").click().run()
    assert not at.exception
    lignes = _lignes_du_pptx(at.session_state["pptx_bytes"])
    chant = [l for l in lignes if l[0] in ("Refrain inventé de bout en bout", "Premier couplet inventé")]
    assert chant == [
        ("Refrain inventé de bout en bout", True),
        ("Premier couplet inventé", False),
        ("Refrain inventé de bout en bout", True),
    ]
