"""Tests de tools/import_chants/extract_docx.py (documents Word fabriqués par les tests)."""

import io
import threading
import zipfile

import pytest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from tests.helpers_import import docx_avec_entrees_en_plus, docx_bytes, docx_partie_remplacee
from tools.import_chants.extract_docx import extract_docx
from tools.import_chants.modeles import UnsupportedFile


def test_un_titre_souligne_suivi_de_paroles_dans_le_meme_paragraphe_donne_deux_lignes():
    lignes = extract_docx(docx_bytes([[("Psaume", "gs"), (" Le fleuve chante la paix.", "g")]]))
    assert [l.texte for l in lignes] == ["Psaume", "Le fleuve chante la paix."]
    assert lignes[0].gras and lignes[0].souligne
    assert lignes[1].gras and not lignes[1].souligne


def test_paragraphes_vides_deviennent_vide_avant_de_la_ligne_suivante():
    lignes = extract_docx(docx_bytes(["premier vers", "second vers", "", "", "troisième vers"]))
    assert [(l.texte, l.vide_avant) for l in lignes] == [
        ("premier vers", False), ("second vers", False), ("troisième vers", True)]


def test_retours_a_la_ligne_et_doubles_espaces_separent_les_vers():
    paragraphe = "le vent du soir se lève sur la ville  les cloches annoncent le jour\nla nuit vient"
    assert [l.texte for l in extract_docx(docx_bytes([paragraphe]))] == [
        "le vent du soir se lève sur la ville", "les cloches annoncent le jour", "la nuit vient"]


def test_titre_souligne_garde_ses_espaces_et_n_est_pas_coupe():
    lignes = extract_docx(docx_bytes([[("Gloire a Dieu", "gs"), ("   ", "g"), ("Recueil Aurore 2", "gs")]]))
    assert [l.texte for l in lignes] == ["Gloire a Dieu   Recueil Aurore 2"]


def test_gras_et_italique_par_ligne():
    lignes = extract_docx(docx_bytes([[("refrain en gras", "g")], [("vers en italique", "i")], "vers simple"]))
    assert [(l.gras, l.italique) for l in lignes] == [(True, False), (False, True), (False, False)]


def test_le_gras_peut_venir_du_style_du_paragraphe():
    lignes = extract_docx(docx_bytes(["texte sans gras sur le run"], style_gras=True))
    assert lignes[0].gras is True


def test_fichier_illisible():
    with pytest.raises(UnsupportedFile) as erreur:
        extract_docx(b"ceci n'est pas un fichier Word")
    assert "Word" in erreur.value.raison


def _bytes(document: Document) -> bytes:
    """Sauvegarde un document Word en bytes."""
    sortie = io.BytesIO()
    document.save(sortie)
    return sortie.getvalue()


def test_runs_dans_insertions_balises_champs_sdt():
    """Runs dans w:ins, w:smartTag, w:fldSimple et w:sdt sont lus."""
    doc = Document()
    p = doc.add_paragraph("avant ")
    p._p.append(parse_xml('<w:ins %s w:id="1" w:author="test"><w:r><w:t>inséré </w:t></w:r></w:ins>' % nsdecls("w")))
    p._p.append(parse_xml('<w:smartTag %s w:uri="x" w:element="y"><w:r><w:t>balise </w:t></w:r></w:smartTag>' % nsdecls("w")))
    p._p.append(parse_xml('<w:fldSimple %s w:instr="PAGE"><w:r><w:t>champ </w:t></w:r></w:fldSimple>' % nsdecls("w")))
    p._p.append(parse_xml('<w:sdt %s><w:sdtContent><w:r><w:t>contrôle</w:t></w:r></w:sdtContent></w:sdt>' % nsdecls("w")))
    lignes = extract_docx(_bytes(doc))
    assert [l.texte for l in lignes] == ["avant inséré balise champ contrôle"]


def test_paragraphe_contenant_uniquement_un_run_insere():
    """Un paragraphe avec seulement un w:ins run n'est pas vide."""
    doc = Document()
    doc.add_paragraph("a")
    p = doc.add_paragraph()
    p._p.append(parse_xml('<w:ins %s w:id="1" w:author="test"><w:r><w:t>b</w:t></w:r></w:ins>' % nsdecls("w")))
    doc.add_paragraph("c")
    lignes = extract_docx(_bytes(doc))
    assert [(l.texte, l.vide_avant) for l in lignes] == [("a", False), ("b", False), ("c", False)]


def test_vide_avant_avec_paragraphes_vides_intermediaires():
    """Paragraphes ["a", "", "b", "c"] → a (False), b (True), c (False)."""
    lignes = extract_docx(docx_bytes(["a", "", "b", "c"]))
    assert [(l.texte, l.vide_avant) for l in lignes] == [("a", False), ("b", True), ("c", False)]


def test_hyperlink_avec_formattage():
    """Texte dans w:hyperlink avec gras est lu et le gras est détecté."""
    doc = Document()
    p = doc.add_paragraph()
    p._p.append(parse_xml('<w:hyperlink %s><w:r><w:rPr><w:b/></w:rPr><w:t>lien gras</w:t></w:r></w:hyperlink>' % nsdecls("w")))
    lignes = extract_docx(_bytes(doc))
    assert [l.texte for l in lignes] == ["lien gras"]
    assert lignes[0].gras is True


def test_style_paragraphe_parent_avec_gras():
    """Style enfant hérité de parent avec gras sur deux niveaux."""
    doc = Document()
    style_parent = doc.styles.add_style("Parent", WD_STYLE_TYPE.PARAGRAPH)
    style_parent.font.bold = True
    style_enfant = doc.styles.add_style("Enfant", WD_STYLE_TYPE.PARAGRAPH)
    style_enfant.base_style = style_parent
    p = doc.add_paragraph(style=style_enfant)
    run = p.add_run("texte")
    lignes = extract_docx(_bytes(doc))
    assert lignes[0].gras is True


def test_style_caractere_avec_gras():
    """Style caractère appliqué à un run avec gras."""
    doc = Document()
    style_char = doc.styles.add_style("GrasChar", WD_STYLE_TYPE.CHARACTER)
    style_char.font.bold = True
    p = doc.add_paragraph()
    run = p.add_run("texte")
    run.style = style_char
    lignes = extract_docx(_bytes(doc))
    assert lignes[0].gras is True


def test_run_bold_false_override_style_bold():
    """Un run avec bold=False override le style bold=True."""
    doc = Document()
    style = doc.styles.add_style("GrasStyle", WD_STYLE_TYPE.PARAGRAPH)
    style.font.bold = True
    p = doc.add_paragraph(style=style)
    run = p.add_run("texte")
    run.bold = False
    lignes = extract_docx(_bytes(doc))
    assert lignes[0].gras is False


def test_tableau_est_ignore():
    """Un tableau dans le document n'est pas lu, pas d'exception."""
    doc = Document()
    doc.add_paragraph("avant")
    table = doc.add_table(1, 1)
    table.rows[0].cells[0].paragraphs[0].add_run("dans le tableau")
    doc.add_paragraph("après")
    lignes = extract_docx(_bytes(doc))
    assert [l.texte for l in lignes] == ["avant", "après"]


def test_xml_corrompu_dans_zip_valide():
    """Un fichier .docx avec document.xml corrompu lève UnsupportedFile."""
    doc = Document()
    doc.add_paragraph("texte")
    data = _bytes(doc)
    with zipfile.ZipFile(io.BytesIO(data), "r") as zin:
        files = {name: zin.read(name) if name != "word/document.xml" else b"<w:document" for name in zin.namelist()}
    sortie = io.BytesIO()
    with zipfile.ZipFile(sortie, "w") as zout:
        for name, content in files.items():
            zout.writestr(name, content)
    with pytest.raises(UnsupportedFile) as erreur:
        extract_docx(sortie.getvalue())
    assert "Word" in erreur.value.raison


def test_style_cyclique_ne_boucle_pas():
    """Un style avec base_style cyclique ne cause pas de boucle infinie."""
    doc = Document()
    style = doc.styles.add_style("Cyclique", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = style
    p = doc.add_paragraph("texte", style=style)
    thread = threading.Thread(target=lambda: extract_docx(_bytes(doc)), daemon=True)
    thread.start()
    thread.join(timeout=5)
    assert not thread.is_alive(), "extract_docx est bloquée sur un style cyclique"


# --- I1. Un .docx abîmé mais bien formé donne UnsupportedFile, jamais une exception brute ---

_W = b"http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _docx_abime(cas: str) -> bytes:
    if cas == "racine_inattendue":
        return docx_partie_remplacee(docx_bytes(["vers"]), "word/document.xml", lambda _: b'<racine xmlns:w="' + _W + b'"/>')
    if cas == "sans_body":
        return docx_partie_remplacee(
            docx_bytes(["vers"]), "word/document.xml", lambda _: b'<w:document xmlns:w="' + _W + b'"/>')
    if cas == "gras_invalide":
        return docx_partie_remplacee(
            docx_bytes([[("vers", "g")]]), "word/document.xml", lambda x: x.replace(b"<w:b/>", b'<w:b w:val="maybe"/>'))
    if cas == "souligne_invalide":
        return docx_partie_remplacee(
            docx_bytes([[("vers", "s")]]), "word/document.xml",
            lambda x: x.replace(b'<w:u w:val="single"/>', b'<w:u w:val="weird"/>'))
    if cas == "styles_racine_inattendue":
        return docx_partie_remplacee(
            docx_bytes(["vers"], style_gras=True), "word/styles.xml", lambda _: b'<racine xmlns:w="' + _W + b'"/>')
    if cas == "ooxml_strict":
        strict = b"http://purl.oclc.org/ooxml/wordprocessingml/main"
        return docx_partie_remplacee(docx_bytes(["vers"]), "word/document.xml", lambda x: x.replace(_W, strict))
    raise AssertionError(cas)


@pytest.mark.parametrize("cas", ["racine_inattendue", "sans_body", "gras_invalide", "souligne_invalide",
                                 "styles_racine_inattendue", "ooxml_strict"])
def test_docx_abime_mais_bien_forme_est_refuse_proprement(cas):
    with pytest.raises(UnsupportedFile) as erreur:
        extract_docx(_docx_abime(cas))
    assert "Word" in erreur.value.raison


# --- I2. Garde anti-bombe de décompression, et résolution des styles mémoïsée ---

MO = 1024 * 1024
ENTREES_D_UN_DOCX = len(zipfile.ZipFile(io.BytesIO(docx_bytes(["x"]))).namelist())


@pytest.mark.parametrize("entrees", [
    {"word/document.xml": bytes(6 * MO)},
    {f"word/media/{k}.bin": b"x" for k in range(1001 - ENTREES_D_UN_DOCX)},
    {f"word/media/{k}.bin": bytes(6 * MO) for k in range(9)},
], ids=["document_xml_de_6_mo", "mille_et_une_entrees", "somme_de_54_mo"])
def test_docx_trop_volumineux_une_fois_decompresse_est_refuse_avant_lecture(entrees):
    with pytest.raises(UnsupportedFile) as erreur:
        extract_docx(docx_avec_entrees_en_plus(entrees))
    assert erreur.value.raison == "Fichier Word trop volumineux une fois décompressé"


def test_docx_dans_les_limites_est_lu():
    data = docx_avec_entrees_en_plus({f"word/media/{k}.bin": b"x" for k in range(1000 - ENTREES_D_UN_DOCX)})
    assert [l.texte for l in extract_docx(data)] == ["un vers inventé"]


def test_la_resolution_des_styles_est_memoisee_par_style_et_attribut(monkeypatch):
    from docx.styles.styles import Styles
    appels = []
    original = Styles.get_by_id
    monkeypatch.setattr(Styles, "get_by_id", lambda self, *a, **k: appels.append(a) or original(self, *a, **k))
    lignes = extract_docx(docx_bytes(["vers inventé"] * 2000, style_gras=True))
    assert len(lignes) == 2000 and all(l.gras for l in lignes)
    assert len(appels) < 40  # une résolution par style et par attribut, pas une par run


# --- Mineur : le souligné peut être un type de trait ---

def test_souligne_double_ou_pointille_est_souligne_et_aucun_soulignement_ne_l_est_pas():
    from docx.enum.text import WD_UNDERLINE
    doc = Document()
    p = doc.add_paragraph()
    for texte, valeur in [("double", WD_UNDERLINE.DOUBLE), ("pointille", WD_UNDERLINE.DOTTED),
                          ("aucun", WD_UNDERLINE.NONE), ("faux", False)]:
        run = p.add_run(texte + "\n")
        run.underline = valeur
    lignes = extract_docx(_bytes(doc))
    assert [(l.texte, l.souligne) for l in lignes] == [
        ("double", True), ("pointille", True), ("aucun", False), ("faux", False)]
