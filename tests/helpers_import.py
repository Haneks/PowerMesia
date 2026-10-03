"""Fabrique de fichiers Word et PDF pour les tests de l'import (textes inventés, rien du corpus)."""

import io

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from tools.import_chants.modeles import Line


def L(texte: str, gras=False, italique=False, souligne=False, vide=False) -> Line:
    """Une Line abrégée, pour tester l'analyse sans passer par un fichier."""
    return Line(texte, gras, italique, souligne, vide)


def docx_bytes(paragraphes: list, style_gras: bool = False) -> bytes:
    """
    Un .docx. Chaque paragraphe est "" (paragraphe vide), une chaîne (texte simple) ou une liste de
    segments (texte, drapeaux) où les drapeaux sont des lettres : g = gras, i = italique, s = souligné.
    style_gras : applique à tous les paragraphes un style de paragraphe en gras (sans gras sur les runs).
    """
    document = Document()
    style = None
    if style_gras:
        style = document.styles.add_style("TitreGras", WD_STYLE_TYPE.PARAGRAPH)
        style.font.bold = True
    for paragraphe in paragraphes:
        segments = [(paragraphe, "")] if isinstance(paragraphe, str) else paragraphe
        p = document.add_paragraph(style=style)
        for texte, drapeaux in segments:
            run = p.add_run(texte)
            if "g" in drapeaux:
                run.bold = True
            if "i" in drapeaux:
                run.italic = True
            if "s" in drapeaux:
                run.underline = True
    sortie = io.BytesIO()
    document.save(sortie)
    return sortie.getvalue()
