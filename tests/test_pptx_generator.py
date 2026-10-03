"""Tests de la génération PPTX (tools/pptx_generator.py) : style, titres, pagination."""

import re

import pytest
from pptx import Presentation
from pptx.util import Pt

from tests.test_slicing import EVANGILE
from tools.pptx_generator import generate_pptx

CHANT = (
    "Qu’il est bon, qu’il est doux,\nQue les frères soient unis,\n"
    "Dans la paix, dans la joie,\nDans l’amour de Jésus-Christ.\n\n"
    "Je veux chanter ton nom, Seigneur,\nTu es mon roc, ma forteresse,\n"
    "Mon bouclier, ma délivrance,\nMon refuge et ma joie."
)

BLOCS = [
    {"type": "lecture", "reference": "Jn 3, 16-18", "intro_lue": "Évangile de Jésus Christ selon saint Jean",
     "contenu": f"<p>{EVANGILE}</p>"},
    {"type": "chant", "titre": "Qu’il est bon", "paroles": CHANT},
    {"type": "message", "titre": "Annonces", "contenu": "Quête pour les travaux de l’église."},
]


@pytest.fixture(scope="module")
def slides(tmp_path_factory):
    out = tmp_path_factory.mktemp("pptx") / "messe.pptx"
    generate_pptx(BLOCS, out)
    prs = Presentation(str(out))
    result = []
    for slide in prs.slides:
        boxes = [sh for sh in slide.shapes if sh.has_text_frame]
        title_box, body_box = boxes[0], boxes[1]
        result.append({"slide": slide, "title": title_box, "body": body_box})
    return result


def _title_text(s):
    return s["title"].text_frame.text


def _body_text(s):
    return s["body"].text_frame.text


def _runs(shape):
    return [r for p in shape.text_frame.paragraphs for r in p.runs]


def test_title_has_x_over_y_pagination_per_block(slides):
    titles = [_title_text(s) for s in slides]
    parsed = [re.fullmatch(r"(.+) - (\d+)/(\d+)", t) for t in titles]
    assert all(parsed), titles

    by_label: dict[str, list[tuple[int, int]]] = {}
    for m in parsed:
        by_label.setdefault(m.group(1), []).append((int(m.group(2)), int(m.group(3))))

    assert set(by_label) == {"Évangile de Jésus Christ selon saint Jean", "Qu’il est bon", "Annonces"}
    for label, pages in by_label.items():
        y = len(pages)
        assert pages == [(x, y) for x in range(1, y + 1)], (label, pages)


def test_single_slide_block_is_numbered_1_over_1(slides):
    assert [_title_text(s) for s in slides if _title_text(s).startswith("Annonces")] == ["Annonces - 1/1"]


def test_body_is_calibri_54_black(slides):
    for s in slides:
        for run in _runs(s["body"]):
            assert run.font.name == "Calibri"
            assert run.font.size == Pt(54)
            assert str(run.font.color.rgb) == "000000"


def test_title_is_calibri_and_black(slides):
    for s in slides:
        for run in _runs(s["title"]):
            assert run.font.name == "Calibri"
            assert str(run.font.color.rgb) == "000000"


def test_body_text_is_at_most_150_chars(slides):
    for s in slides:
        assert len(_body_text(s)) <= 150


def test_html_is_stripped_and_text_preserved(slides):
    body = " ".join(_body_text(s) for s in slides if s["title"].text_frame.text.startswith("Évangile"))
    assert "<" not in body
    assert re.sub(r"\s+", " ", body) == re.sub(r"\s+", " ", EVANGILE)


def test_chant_keeps_line_breaks(slides):
    chant_bodies = [_body_text(s) for s in slides if _title_text(s).startswith("Qu’il est bon")]
    assert any("\n" in b or "\v" in b for b in chant_bodies)


def test_background_is_light_so_black_text_is_readable(slides):
    for s in slides:
        rgb = s["slide"].background.fill.fore_color.rgb
        assert min(rgb[0], rgb[1], rgb[2]) >= 0xE0, rgb


def test_html_entities_are_decoded():
    from tools.pptx_generator import _strip_html

    assert _strip_html("<p>L&#39;amour&nbsp;: Dieu &amp; nous</p>") == "L'amour\xa0: Dieu & nous"
